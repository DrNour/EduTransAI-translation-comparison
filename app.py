import re
import unicodedata
from difflib import SequenceMatcher
from hashlib import blake2b
from io import BytesIO
from typing import Any

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
import streamlit as st
from docx import Document
from nltk.translate.bleu_score import SmoothingFunction, sentence_bleu
from sentence_transformers import SentenceTransformer

# Optional lexical metric: sacrebleu CHRF (preferred)
try:
    from sacrebleu.metrics import CHRF

    _CHRF = CHRF(word_order=0)
except Exception:
    _CHRF = None

# Optional semantic metric: BERTScore
try:
    from bert_score import score as bertscore_score

    _HAS_BERTSCORE = True
except Exception:
    _HAS_BERTSCORE = False

st.set_page_config(page_title="EduTransAI - Translation Assessment", layout="wide")
st.title("EduTransAI - Translation Comparison & Student Assessment")


# ===========================
# Sidebar controls
# ===========================
with st.sidebar:
    st.header("Settings")
    model_name = st.selectbox(
        "Embedding model",
        options=[
            "all-MiniLM-L6-v2",
            "paraphrase-multilingual-MiniLM-L12-v2",
        ],
        index=0,
        help="Choose the embedding model for semantic similarity.",
    )
    semantic_weight = st.slider("Semantic weight (for Hybrid)", 0.0, 1.0, 0.65, 0.05)
    lexical_weight = 1.0 - semantic_weight

    sem_choice = st.selectbox(
        "Semantic metric",
        options=["Cosine", "BERTScore", "Cosine+BERTScore"],
        index=0,
        help="BERTScore is optional and falls back automatically if unavailable.",
    )

    use_chrf = st.checkbox(
        "Use CHRF for lexical overlap (fallback to BLEU if unavailable)",
        value=True,
    )

    st.markdown("---")
    st.subheader("Stability & Thresholds")
    fluency_floor = st.number_input("Fluency floor to trust Hybrid", 1.0, 5.0, 3.0, 0.1)
    style_floor = st.number_input("Style floor to trust Hybrid", 1.0, 5.0, 3.0, 0.1)
    consistency_tolerance = st.slider("Consistency tolerance (|Cosine - BERTScore|)", 0.0, 1.0, 0.20, 0.01)

    st.markdown("---")
    st.subheader("Evaluation Strategy")
    rebalance_sem = st.checkbox(
        "Emphasize semantic + fluency in Hybrid (rebalance)",
        value=True,
        help="Hybrid ~= 0.8*Semantic + 0.2*Lexical, stabilized by Fluency/Style floors.",
    )
    use_composite = st.checkbox(
        "Report composite indices (SQI & LI)",
        value=True,
        help="SQI = 0.7*Semantic + 0.3*(Fluency/5). LI = Lexical.",
    )
    gateA_threshold = st.slider("Gate A - Semantic threshold", 0.50, 0.99, 0.80, 0.01)
    gateB_lex_threshold = st.slider("Gate B - Lexical threshold", 0.00, 0.80, 0.20, 0.01)
    gateB_flu_threshold = st.slider("Gate B - Fluency threshold", 1.0, 5.0, 3.0, 0.1)

    st.markdown("---")
    st.subheader("Paraphrase vs Drift")
    paraphrase_sem_hi = st.slider("Paraphrase: semantic high >=", 0.70, 0.99, 0.85, 0.01)
    drift_sem_lo = st.slider("Meaning drift: semantic low <", 0.40, 0.95, 0.70, 0.01)
    low_lexical_for_paraphrase = st.slider("Paraphrase lexical threshold <", 0.0, 0.6, 0.20, 0.01)

    st.markdown("---")
    st.subheader("Metadata (optional)")
    st.caption("Add metadata for per-domain diagnostics and export.")
    meta_mode = st.radio("Metadata source", ["None", "Use existing columns", "Set constants"], index=0)

    language_val = domain_val = genre_val = None
    language_col = domain_col = genre_col = None

    if meta_mode == "Use existing columns":
        st.caption("You will be able to select columns after uploading.")
    elif meta_mode == "Set constants":
        language_val = st.text_input("Language (constant for all rows)", value="")
        domain_val = st.text_input("Domain (constant for all rows)", value="")
        genre_val = st.text_input("Genre (constant for all rows)", value="")


# ===========================
# Cached loaders
# ===========================
@st.cache_resource(show_spinner=True)
def load_model(name: str):
    return SentenceTransformer(name)


@st.cache_data(show_spinner=False)
def batch_encode_unique(texts: list[str], model_name_for_cache: str):
    if not texts:
        return {}
    vecs = load_model(model_name_for_cache).encode(
        texts,
        batch_size=128,
        show_progress_bar=False,
        normalize_embeddings=True,
    )
    return {text_key(t): v for t, v in zip(texts, vecs)}


# ===========================
# Helper functions
# ===========================
word_or_punct = re.compile(r"[A-Za-zÀ-ÖØ-öø-ÿ0-9]+|[^\w\s]", re.UNICODE)


def normalize_text(t: str) -> str:
    t = unicodedata.normalize("NFKC", str(t)).strip().lower()
    t = re.sub(r"[“”]", '"', t)
    t = re.sub(r"[‘’]", "'", t)
    t = re.sub(r"\s+", " ", t)
    return t


def simple_tokenize(text: str):
    return word_or_punct.findall(text)


def text_key(t: str) -> str:
    return blake2b(t.encode("utf-8"), digest_size=12).hexdigest()


def get_vec(t: str, cache: dict):
    return cache.get(text_key(t))


def chrf_score(ref: str, hyp: str) -> float:
    if _CHRF is None:
        raise NameError("CHRF not available")
    return _CHRF.sentence_score(hyp, [ref]).score / 100.0


def length_ratio_penalty(ref: str, hyp: str) -> float:
    r = max(1e-6, len(hyp.split())) / max(1e-6, len(ref.split()))
    return float(np.exp(-abs(np.log(r))))


def compute_style_features(t: str) -> dict:
    t = normalize_text(t)
    t = re.sub(r"\s+", " ", t.strip())
    sents = max(1, len(re.findall(r"[.!?]", t)))
    tokens = t.split()
    n_tok = max(1, len(tokens))

    avg_sent_len = n_tok / sents
    long_tok_ratio = sum(len(w) > 24 for w in tokens) / n_tok
    repeat_punct = 1 if re.search(r"([!?.,])\1{1,}", t) else 0
    double_space = 1 if re.search(r"\s{2,}", t) else 0
    trailing_space = 1 if re.search(r"\s$", t) else 0
    ttr = len(set(tokens)) / n_tok

    score = 5.0
    score -= min(2.0, abs(avg_sent_len - 18) / 18) * 1.2
    score -= long_tok_ratio * 2.0
    score -= (repeat_punct + double_space + trailing_space) * 0.5
    score -= max(0, 0.35 - ttr) * 2.0
    score = float(np.clip(score, 1.0, 5.0))

    return {
        "avg_sent_len": avg_sent_len,
        "long_tok_ratio": long_tok_ratio,
        "repeat_punct": repeat_punct,
        "double_space": double_space,
        "trailing_space": trailing_space,
        "ttr": round(ttr, 3),
        "style_score": round(score, 2),
    }


def fluency_score(text: str) -> float:
    t = normalize_text(text)
    if not t:
        return 1.0
    tokens = t.split()
    long_token_ratio = sum(len(w) > 24 for w in tokens) / max(1, len(tokens))
    punct_endings = len(re.findall(r"[.!?]", t))
    sentences = max(1, punct_endings)
    avg_sent_len = len(tokens) / sentences
    repeats = len(re.findall(r"([!?.,])\1{1,}", t))
    weird_ws = 1 if re.search(r"\s{2,}", t) else 0

    score = 5.0
    score -= min(2.0, abs(avg_sent_len - 18) / 18) * 1.0
    score -= long_token_ratio * 2.0
    score -= repeats * 0.3
    score -= weird_ws * 0.5
    return round(float(np.clip(score, 1.0, 5.0)), 2)


def token_diff(a: str, b: str) -> str:
    a_t, b_t = a.split(), b.split()
    sm = SequenceMatcher(None, a_t, b_t)
    parts = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            parts.extend(b_t[j1:j2])
        elif tag == "replace":
            if i1 != i2:
                parts.append(f"<span style='background:#ffe6e6;text-decoration:line-through'>{' '.join(a_t[i1:i2])}</span>")
            if j1 != j2:
                parts.append(f"<span style='background:#e6ffe6;'>{' '.join(b_t[j1:j2])}</span>")
        elif tag == "delete":
            parts.append(f"<span style='background:#ffe6e6;text-decoration:line-through'>{' '.join(a_t[i1:i2])}</span>")
        elif tag == "insert":
            parts.append(f"<span style='background:#e6ffe6;'>{' '.join(b_t[j1:j2])}</span>")
    return " ".join(parts)


def _safe_clip01(x):
    try:
        return float(np.clip(x, 0.0, 1.0))
    except Exception:
        return 0.0


def compute_semantic_signals(ref: str, hyp: str, embed_cache: dict, sem_choice: str, bert_cache: dict | None = None):
    ref_n, hyp_n = normalize_text(ref), normalize_text(hyp)
    v_ref, v_hyp = get_vec(ref_n, embed_cache), get_vec(hyp_n, embed_cache)
    cosine = float(np.dot(v_ref, v_hyp)) if (v_ref is not None and v_hyp is not None) else None

    bert_f1 = None
    if _HAS_BERTSCORE and sem_choice in ("BERTScore", "Cosine+BERTScore"):
        cache_key = text_key(ref_n + " ||| " + hyp_n)
        if bert_cache is not None and cache_key in bert_cache:
            bert_f1 = bert_cache[cache_key]
        else:
            try:
                _, _, f1 = bertscore_score([hyp_n], [ref_n], verbose=False, rescale_with_baseline=False)
                bert_f1 = _safe_clip01(float(f1[0].item()))
                if bert_cache is not None:
                    bert_cache[cache_key] = bert_f1
            except Exception:
                bert_f1 = None

    return {"cosine": cosine, "bertscore_f1": bert_f1}


def semantic_accuracy_score(
    ref: str,
    hyp: str,
    embed_cache: dict,
    w_sem: float,
    w_lex: float,
    prefer_chrf: bool,
    sem_choice: str,
    fluency: float | None = None,
    style: float | None = None,
    fluency_floor: float = 3.0,
    style_floor: float = 3.0,
    consistency_tolerance: float = 0.20,
    bert_cache: dict | None = None,
):
    ref_n, hyp_n = normalize_text(ref), normalize_text(hyp)

    if prefer_chrf and _CHRF is not None:
        lexical = chrf_score(ref_n, hyp_n)
        used_lex = "CHRF"
    else:
        lexical = float(
            sentence_bleu(
                [simple_tokenize(ref_n)],
                simple_tokenize(hyp_n),
                smoothing_function=SmoothingFunction().method4,
            )
        )
        used_lex = "BLEU"

    sigs = compute_semantic_signals(ref_n, hyp_n, embed_cache, sem_choice, bert_cache=bert_cache)
    cosine = sigs.get("cosine")
    bert_f1 = sigs.get("bertscore_f1")

    if sem_choice == "BERTScore" and bert_f1 is not None:
        sem_val, used_sem = bert_f1, "BERTScore"
    elif sem_choice == "Cosine+BERTScore" and (cosine is not None or bert_f1 is not None):
        vals = [v for v in [cosine, bert_f1] if v is not None]
        sem_val, used_sem = float(np.mean(vals)), "Cosine+BERTScore"
    else:
        sem_val, used_sem = (cosine if cosine is not None else 0.0), "Cosine"

    penalty = length_ratio_penalty(ref_n, hyp_n)
    hybrid = _safe_clip01(w_sem * sem_val + w_lex * lexical) * penalty

    if fluency is not None and fluency < fluency_floor:
        hybrid *= max(0.1, float(fluency) / float(fluency_floor))
    if style is not None and style < style_floor:
        hybrid *= max(0.1, float(style) / float(style_floor))

    if cosine is not None and bert_f1 is not None:
        gap = abs(cosine - bert_f1)
        if gap > consistency_tolerance:
            hybrid *= 1.0 - min(0.30, (gap - consistency_tolerance))

    extras = {
        "cosine": None if cosine is None else round(cosine, 3),
        "bertscore_f1": None if bert_f1 is None else round(bert_f1, 3),
        "length_penalty": round(penalty, 3),
        "lexical_metric": used_lex,
    }
    return round(sem_val, 3), used_sem, round(lexical, 3), round(hybrid, 3), extras


def dynamic_thresholds(ref_len_tokens: int):
    acc = 0.60
    lex = 0.50
    if ref_len_tokens >= 25:
        acc -= 0.05
        lex -= 0.05
    elif ref_len_tokens <= 6:
        acc += 0.05
        lex += 0.05
    return acc, lex


def classify_semantic_deviation(
    lexical: float,
    cosine: float | None,
    bert_f1: float | None,
    paraphrase_sem_hi: float = 0.85,
    low_lex: float = 0.20,
    drift_sem_lo: float = 0.70,
) -> str | None:
    sem_signal = bert_f1 if bert_f1 is not None else cosine
    if sem_signal is None:
        return None
    if sem_signal >= paraphrase_sem_hi and lexical < low_lex:
        return "Paraphrase"
    if sem_signal < drift_sem_lo:
        return "Meaning Drift"
    return None


LABEL_CANON = {
    "fluency": ["fluency/grammar", "grammar", "fluency issues"],
    "style": ["style", "formatting", "layout"],
    "semantic": ["semantic deviation", "meaning error", "mismatch"],
    "lexical": ["low lexical overlap", "lexical low", "overlap low"],
    "length_short": ["too short", "short"],
    "length_long": ["too long / verbosity", "too long", "verbosity"],
    "paraphrase": ["paraphrase", "lexically diverse paraphrase"],
    "meaning_drift": ["meaning drift", "semantic shift"],
}
_LABEL_LOOKUP = {k: k for k in LABEL_CANON}
for k, aliases in LABEL_CANON.items():
    for alias in aliases:
        _LABEL_LOOKUP[alias.lower()] = k


def normalize_labels(raw_list):
    norm = []
    for raw in (raw_list or []):
        canon = _LABEL_LOOKUP.get(str(raw).strip().lower())
        norm.append(canon if canon else f"other:{raw}")
    seen, out = set(), []
    for item in norm:
        if item not in seen:
            out.append(item)
            seen.add(item)
    return out


KNOWN_SUFFIXES = sorted(
    [
        "Low_Hybrid_Flag_Adaptive",
        "Low_Lexical_Flag_Adaptive",
        "Low_Hybrid_Flag_Global",
        "Low_Lexical_Flag_Global",
        "Low_Fluency_Flag_Global",
        "Low_Style_Flag_Global",
        "GateA_SemanticOK",
        "SemanticMetric",
        "LexicalMetric",
        "BERTScoreF1",
        "GateB_Flag",
        "ErrorsNorm",
        "Accuracy",
        "Semantic",
        "Fluency",
        "Lexical",
        "Cosine",
        "Style",
        "Errors",
        "SQI",
        "LI",
    ],
    key=len,
    reverse=True,
)


def _pretty_metric_suffix(suffix: str) -> str:
    mapping_exact = {
        "Accuracy": " (Hybrid Accuracy)",
        "Lexical": " (Lexical)",
        "Semantic": " (Semantic Score)",
        "Cosine": " (Semantic Cosine)",
        "BERTScoreF1": " (BERTScore F1)",
        "Fluency": " (Fluency)",
        "Style": " (Style)",
        "SQI": " (Semantic Quality Index)",
        "LI": " (Literalness Index)",
        "Errors": " (Error Categories - Raw)",
        "ErrorsNorm": " (Error Labels - Canonical)",
        "SemanticMetric": " (Semantic Metric)",
        "LexicalMetric": " (Lexical Metric)",
        "GateA_SemanticOK": " (Gate A - Semantic OK)",
        "GateB_Flag": " (Gate B - Needs Review)",
        "Low_Hybrid_Flag_Global": " Low Hybrid Flag (Global)",
        "Low_Hybrid_Flag_Adaptive": " Low Hybrid Flag (Adaptive)",
        "Low_Lexical_Flag_Global": " Low Lexical Flag (Global)",
        "Low_Lexical_Flag_Adaptive": " Low Lexical Flag (Adaptive)",
        "Low_Fluency_Flag_Global": " Low Fluency Flag (Global)",
        "Low_Style_Flag_Global": " Low Style Flag (Global)",
    }
    return mapping_exact.get(suffix, " " + suffix.replace("_", " "))


def humanize_export_columns(cols) -> list:
    out = []
    seen = set()
    for col in cols:
        new = col
        if isinstance(col, str) and col in {"Language", "Domain", "Genre", "Best_Translation"}:
            new = col
        elif isinstance(col, str):
            for suffix in KNOWN_SUFFIXES:
                token = f"_{suffix}"
                if col.endswith(token):
                    base = col[: -len(token)]
                    new = base.replace("_vs_", " vs ") + _pretty_metric_suffix(suffix)
                    break
        candidate = new
        k = 2
        while candidate in seen:
            candidate = f"{new} #{k}"
            k += 1
        seen.add(candidate)
        out.append(candidate)
    return out


def _col_exists(df, c):
    return c in df.columns


def _safe_num(x):
    try:
        return float(x)
    except Exception:
        return np.nan


def build_analysis_units(mode: str, translation_cols: list[str], source_col: str | None = None):
    units = []
    if mode == "Pairwise Comparison":
        for left in translation_cols:
            for right in translation_cols:
                if left == right:
                    continue
                units.append(
                    {
                        "prefix": f"{left}_vs_{right}",
                        "label": f"{left} vs {right}",
                        "translation_col": left,
                        "reference_col": right,
                    }
                )
    else:
        for col in translation_cols:
            units.append(
                {
                    "prefix": col,
                    "label": col,
                    "translation_col": col,
                    "reference_col": source_col if mode == "Reference-based" else None,
                }
            )
    return units


def make_example_record(i: int, unit: dict, res_df: pd.DataFrame, df_orig: pd.DataFrame, mode: str) -> dict:
    prefix = unit["prefix"]
    rec = {"Row": i, "Student": unit["label"]}
    for suffix in ["Accuracy", "Semantic", "Cosine", "BERTScoreF1", "Lexical", "Fluency", "Style", "ErrorsNorm"]:
        col = f"{prefix}_{suffix}"
        if _col_exists(res_df, col):
            rec[suffix] = res_df.at[i, col] if suffix == "ErrorsNorm" else _safe_num(res_df.at[i, col])

    ref_col = unit.get("reference_col")
    trans_col = unit.get("translation_col")
    rec["Source"] = str(df_orig.at[i, ref_col]) if ref_col and ref_col in df_orig.columns else ""
    rec["Translation"] = str(df_orig.at[i, trans_col]) if trans_col and trans_col in df_orig.columns else ""
    return rec


def gather_issue_examples(
    res_df: pd.DataFrame,
    df_orig: pd.DataFrame,
    analysis_units: list[dict],
    mode: str,
    *,
    top_n: int = 10,
    thresholds: dict | None = None,
    consistency_tolerance: float = 0.20,
    paraphrase_sem_hi: float = 0.85,
    low_lexical_for_paraphrase: float = 0.20,
    drift_sem_lo: float = 0.70,
):
    thresholds = thresholds or {}
    lex_low = thresholds.get("LexLow", 0.15)
    acc_low = thresholds.get("AccLow", 0.50)
    flu_low = thresholds.get("FluLow", 3.0)
    style_low = thresholds.get("StyleLow", 3.0)

    buckets = {
        "Low Accuracy / Semantic": [],
        "Low Lexical Overlap": [],
        "Low Fluency": [],
        "Low Style": [],
        "Paraphrase (Sem High / Lex Low)": [],
        "Meaning Drift (Sem Low)": [],
        "Metric Disagreement (Cos vs BERTScore)": [],
    }

    for i in range(len(res_df)):
        for unit in analysis_units:
            prefix = unit["prefix"]
            sem_v = _safe_num(res_df.at[i, f"{prefix}_Semantic"]) if _col_exists(res_df, f"{prefix}_Semantic") else np.nan
            acc_v = _safe_num(res_df.at[i, f"{prefix}_Accuracy"]) if _col_exists(res_df, f"{prefix}_Accuracy") else np.nan
            lex_v = _safe_num(res_df.at[i, f"{prefix}_Lexical"]) if _col_exists(res_df, f"{prefix}_Lexical") else np.nan
            flu_v = _safe_num(res_df.at[i, f"{prefix}_Fluency"]) if _col_exists(res_df, f"{prefix}_Fluency") else np.nan
            sty_v = _safe_num(res_df.at[i, f"{prefix}_Style"]) if _col_exists(res_df, f"{prefix}_Style") else np.nan
            cos_v = _safe_num(res_df.at[i, f"{prefix}_Cosine"]) if _col_exists(res_df, f"{prefix}_Cosine") else np.nan
            bs_v = _safe_num(res_df.at[i, f"{prefix}_BERTScoreF1"]) if _col_exists(res_df, f"{prefix}_BERTScoreF1") else np.nan

            rec = make_example_record(i, unit, res_df, df_orig, mode)

            if pd.notna(acc_v) and acc_v < acc_low:
                rec1 = dict(rec)
                rec1["SortScore"] = acc_v
                buckets["Low Accuracy / Semantic"].append(rec1)

            if pd.notna(lex_v) and lex_v < lex_low:
                rec2 = dict(rec)
                rec2["SortScore"] = lex_v
                buckets["Low Lexical Overlap"].append(rec2)

            if pd.notna(flu_v) and flu_v < flu_low:
                rec3 = dict(rec)
                rec3["SortScore"] = flu_v
                buckets["Low Fluency"].append(rec3)

            if pd.notna(sty_v) and sty_v < style_low:
                rec4 = dict(rec)
                rec4["SortScore"] = sty_v
                buckets["Low Style"].append(rec4)

            sem_signal = sem_v if pd.notna(sem_v) else (cos_v if pd.notna(cos_v) else np.nan)
            if pd.notna(sem_signal) and pd.notna(lex_v):
                if sem_signal >= paraphrase_sem_hi and lex_v < low_lexical_for_paraphrase:
                    rec5 = dict(rec)
                    rec5["SortScore"] = -sem_signal
                    buckets["Paraphrase (Sem High / Lex Low)"].append(rec5)
                if sem_signal < drift_sem_lo:
                    rec6 = dict(rec)
                    rec6["SortScore"] = sem_signal
                    buckets["Meaning Drift (Sem Low)"].append(rec6)

            if pd.notna(cos_v) and pd.notna(bs_v) and abs(cos_v - bs_v) > consistency_tolerance:
                rec7 = dict(rec)
                rec7["SortScore"] = -abs(cos_v - bs_v)
                buckets["Metric Disagreement (Cos vs BERTScore)"].append(rec7)

    out = {}
    for label, rows in buckets.items():
        if not rows:
            out[label] = pd.DataFrame()
            continue
        dfk = pd.DataFrame(rows).sort_values("SortScore", ascending=True, na_position="last")
        out[label] = dfk.head(top_n).drop(columns=["SortScore"], errors="ignore")

    combined_parts = [v.assign(Issue=k) for k, v in out.items() if not v.empty]
    out["Combined"] = pd.concat(combined_parts, axis=0, ignore_index=True) if combined_parts else pd.DataFrame()
    return out


def examples_to_docx(ex_df: pd.DataFrame, title: str = "Examples") -> bytes:
    doc = Document()
    doc.add_heading(title, level=1)
    for _, row in ex_df.iterrows():
        doc.add_paragraph(f"Row {int(row['Row'])} - {str(row.get('Student', ''))}")
        if "Issue" in ex_df.columns:
            doc.add_paragraph(f"Issue: {str(row.get('Issue', ''))}")
        metric_line = []
        for k in ["Accuracy", "Semantic", "Cosine", "BERTScoreF1", "Lexical", "Fluency", "Style"]:
            if k in ex_df.columns and pd.notna(row.get(k, np.nan)):
                try:
                    metric_line.append(f"{k}={float(row.get(k)):.3f}")
                except Exception:
                    pass
        if metric_line:
            doc.add_paragraph(" | ".join(metric_line))
        if "Source" in ex_df.columns and str(row.get("Source", "")).strip():
            doc.add_paragraph("Source:")
            doc.add_paragraph(str(row.get("Source", "")))
        doc.add_paragraph("Translation:")
        doc.add_paragraph(str(row.get("Translation", "")))
        doc.add_paragraph("--------------------")
    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


def safe_q(series: pd.Series, q: float, default: float):
    series = pd.to_numeric(series, errors="coerce").dropna()
    if series.empty:
        return default
    try:
        return float(series.quantile(q))
    except Exception:
        return default


def collect_metric_series(res_df: pd.DataFrame, suffix: str) -> pd.Series:
    cols = [c for c in res_df.columns if c.endswith(suffix)]
    if not cols:
        return pd.Series(dtype=float)
    return pd.concat([res_df[c] for c in cols], axis=0)


def iqr_outlier_share(series: pd.Series):
    series = pd.to_numeric(series, errors="coerce").dropna()
    if series.empty:
        return np.nan
    q1, q3 = series.quantile(0.25), series.quantile(0.75)
    iqr = q3 - q1
    low, high = q1 - 1.5 * iqr, q3 + 1.5 * iqr
    return float(((series < low) | (series > high)).mean() * 100)


def show_metric_boxplot(data: pd.DataFrame, cols: list[str], metric: str):
    fig, ax = plt.subplots(figsize=(10, 4))
    sns.boxplot(data=data[cols], ax=ax)
    ax.set_ylabel(metric)
    ax.set_title(f"{metric} Distribution Across Students / Translations")
    plt.xticks(rotation=25, ha="right")
    st.pyplot(fig)
    plt.close(fig)


# ===========================
# File upload
# ===========================
uploaded_file = st.file_uploader(
    "Upload CSV, Excel, or Word file (translations)",
    type=["csv", "xlsx", "xls", "docx"],
)

if not uploaded_file:
    st.info("Upload CSV, Excel, or Word file to begin analysis.")
    st.stop()

try:
    file_name = uploaded_file.name.lower()
    if file_name.endswith(".csv"):
        df = pd.read_csv(uploaded_file, encoding="utf-8", na_filter=False)
    elif file_name.endswith((".xlsx", ".xls")):
        df = pd.read_excel(uploaded_file, na_filter=False)
    elif file_name.endswith(".docx"):
        doc = Document(uploaded_file)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        df = pd.DataFrame({"Text": paragraphs})
    else:
        st.error("Unsupported file type.")
        st.stop()

    df.columns = pd.Index([str(c).strip() if str(c).strip() else "Unnamed" for c in df.columns])
    df = df.fillna("")

    st.subheader("Preview of Uploaded Data")
    st.dataframe(df.head())

    df_norm = df.apply(lambda col: col.map(normalize_text))

    mode = st.radio(
        "Select assessment mode:",
        ["Reference-based", "Pairwise Comparison", "Standalone Student Assessment"],
    )

    source_col = None
    translation_cols = []

    if mode == "Reference-based" and len(df.columns) > 1:
        source_col = st.selectbox("Reference / Source Column", list(df.columns))
        translation_cols = st.multiselect(
            "Translations / Student Submissions",
            [c for c in df.columns if c != source_col],
        )
    else:
        translation_cols = st.multiselect("Translations / Student Submissions", list(df.columns))

    if meta_mode == "Use existing columns":
        language_col = st.selectbox("Language column (optional)", ["<none>"] + list(df.columns), index=0)
        domain_col = st.selectbox("Domain column (optional)", ["<none>"] + list(df.columns), index=0)
        genre_col = st.selectbox("Genre column (optional)", ["<none>"] + list(df.columns), index=0)
        language_col = None if language_col == "<none>" else language_col
        domain_col = None if domain_col == "<none>" else domain_col
        genre_col = None if genre_col == "<none>" else genre_col

    if sem_choice in {"BERTScore", "Cosine+BERTScore"} and not _HAS_BERTSCORE:
        st.warning("BERTScore is not installed in this environment. The app will fall back to cosine similarity.")

    if not translation_cols:
        st.warning("Select at least one translation column.")
        st.stop()

    if not st.button("Run Analysis"):
        st.stop()

    st.subheader("Analysis Results")

    use_semantic = mode in {"Reference-based", "Pairwise Comparison"}
    embed_cache = {}
    bert_cache: dict[str, Any] = {}

    if use_semantic:
        all_texts = []
        if source_col:
            all_texts.extend(df_norm[source_col].astype(str).tolist())
        for col in translation_cols:
            all_texts.extend(df_norm[col].astype(str).tolist())
        all_texts = list({t for t in all_texts if t.strip()})
        embed_cache = batch_encode_unique(all_texts, model_name)

    ref_lens = None
    if mode == "Reference-based" and source_col:
        ref_lens = df_norm[source_col].apply(lambda t: len(str(t).split())).tolist()

    analysis_units = build_analysis_units(mode, translation_cols, source_col=source_col)
    progress = st.progress(0.0)

    results = []
    per_row_meta = []

    for idx, row in df.iterrows():
        row_result = {}

        lang_val = language_val if meta_mode == "Set constants" else (row.get(language_col, "") if language_col else "")
        dom_val = domain_val if meta_mode == "Set constants" else (row.get(domain_col, "") if domain_col else "")
        gen_val = genre_val if meta_mode == "Set constants" else (row.get(genre_col, "") if genre_col else "")
        per_row_meta.append({"Language": lang_val or "", "Domain": dom_val or "", "Genre": gen_val or ""})

        for unit in analysis_units:
            prefix = unit["prefix"]
            trans_col = unit["translation_col"]
            trans_text = str(row[trans_col])
            trans_text_norm = str(df_norm.iloc[idx][trans_col])

            flu = fluency_score(trans_text)
            style_feats = compute_style_features(trans_text)
            style = float(style_feats["style_score"])

            if mode == "Reference-based":
                source_text = str(row[source_col])
                source_text_norm = str(df_norm.iloc[idx][source_col])

                sem_sig, used_sem, lexical, hybrid, extras = semantic_accuracy_score(
                    source_text_norm,
                    trans_text_norm,
                    embed_cache,
                    semantic_weight,
                    lexical_weight,
                    prefer_chrf=use_chrf,
                    sem_choice=sem_choice,
                    fluency=flu,
                    style=style,
                    fluency_floor=fluency_floor,
                    style_floor=style_floor,
                    consistency_tolerance=consistency_tolerance,
                    bert_cache=bert_cache,
                )

                acc_thr, lex_thr = dynamic_thresholds(ref_lens[idx])
                errs_raw = []
                if flu < 3:
                    errs_raw.append("Fluency/Grammar")
                if style < 3:
                    errs_raw.append("Style")

                sem_tag = classify_semantic_deviation(
                    lexical=lexical,
                    cosine=extras.get("cosine"),
                    bert_f1=extras.get("bertscore_f1"),
                    paraphrase_sem_hi=paraphrase_sem_hi,
                    low_lex=low_lexical_for_paraphrase,
                    drift_sem_lo=drift_sem_lo,
                )
                if sem_tag == "Paraphrase":
                    errs_raw.append("Paraphrase")
                elif sem_tag == "Meaning Drift":
                    errs_raw.append("Semantic Deviation")

                if hybrid < acc_thr:
                    errs_raw.append("Semantic Deviation")
                if lexical < lex_thr:
                    errs_raw.append("Low Lexical Overlap")

                n_words = len(trans_text_norm.split())
                if n_words < 3:
                    errs_raw.append("Too Short")
                elif n_words > 60:
                    errs_raw.append("Too Long / Verbosity")

                errs_norm = normalize_labels(errs_raw)

                if rebalance_sem:
                    lp = extras.get("length_penalty", 1.0) or 1.0
                    hyb_reb = float(np.clip(0.8 * (sem_sig or 0.0) + 0.2 * (lexical or 0.0), 0, 1)) * lp
                    if flu < fluency_floor:
                        hyb_reb *= max(0.1, float(flu) / float(fluency_floor))
                    if style < style_floor:
                        hyb_reb *= max(0.1, float(style) / float(style_floor))
                    hybrid = round(hyb_reb, 3)

                sqi = round(float(np.clip(0.7 * (sem_sig or 0.0) + 0.3 * (flu / 5.0), 0, 1)), 3) if use_composite else None
                li = round(float(np.clip(lexical or 0.0, 0, 1)), 3) if use_composite else None
                gateA_pass = bool((sem_sig or 0.0) >= gateA_threshold)
                gateB_flag = bool((not gateA_pass) and (((lexical or 0.0) < gateB_lex_threshold) or (flu < gateB_flu_threshold)))

                row_result.update(
                    {
                        f"{prefix}_LexicalMetric": extras.get("lexical_metric"),
                        f"{prefix}_Lexical": lexical,
                        f"{prefix}_SemanticMetric": used_sem,
                        f"{prefix}_Semantic": sem_sig,
                        f"{prefix}_Cosine": extras.get("cosine"),
                        f"{prefix}_BERTScoreF1": extras.get("bertscore_f1"),
                        f"{prefix}_Accuracy": hybrid,
                        f"{prefix}_Fluency": flu,
                        f"{prefix}_Style": style,
                        f"{prefix}_SQI": sqi,
                        f"{prefix}_LI": li,
                        f"{prefix}_GateA_SemanticOK": "OK" if gateA_pass else "",
                        f"{prefix}_GateB_Flag": "Review" if gateB_flag else "",
                        f"{prefix}_Errors": ", ".join(sorted(set(errs_raw))) if errs_raw else "None",
                        f"{prefix}_ErrorsNorm": ",".join(errs_norm) if errs_norm else "",
                    }
                )

            elif mode == "Pairwise Comparison":
                other_col = unit["reference_col"]
                other_text_norm = str(df_norm.iloc[idx][other_col])

                sem_sig, used_sem, lexical, hybrid, extras = semantic_accuracy_score(
                    other_text_norm,
                    trans_text_norm,
                    embed_cache,
                    semantic_weight,
                    lexical_weight,
                    prefer_chrf=use_chrf,
                    sem_choice=sem_choice,
                    fluency=flu,
                    style=style,
                    fluency_floor=fluency_floor,
                    style_floor=style_floor,
                    consistency_tolerance=consistency_tolerance,
                    bert_cache=bert_cache,
                )

                errs_raw = []
                if flu < 3:
                    errs_raw.append("Fluency/Grammar")
                if style < 3:
                    errs_raw.append("Style")
                n_words = len(trans_text_norm.split())
                if n_words < 3:
                    errs_raw.append("Too Short")
                elif n_words > 60:
                    errs_raw.append("Too Long / Verbosity")
                errs_norm = normalize_labels(errs_raw)

                if rebalance_sem:
                    lp = extras.get("length_penalty", 1.0) or 1.0
                    hyb_reb = float(np.clip(0.8 * (sem_sig or 0.0) + 0.2 * (lexical or 0.0), 0, 1)) * lp
                    if flu < fluency_floor:
                        hyb_reb *= max(0.1, float(flu) / float(fluency_floor))
                    if style < style_floor:
                        hyb_reb *= max(0.1, float(style) / float(style_floor))
                    hybrid = round(hyb_reb, 3)

                sqi = round(float(np.clip(0.7 * (sem_sig or 0.0) + 0.3 * (flu / 5.0), 0, 1)), 3) if use_composite else None
                li = round(float(np.clip(lexical or 0.0, 0, 1)), 3) if use_composite else None
                gateA_pass = bool((sem_sig or 0.0) >= gateA_threshold)
                gateB_flag = bool((not gateA_pass) and (((lexical or 0.0) < gateB_lex_threshold) or (flu < gateB_flu_threshold)))

                row_result.update(
                    {
                        f"{prefix}_LexicalMetric": extras.get("lexical_metric"),
                        f"{prefix}_Lexical": lexical,
                        f"{prefix}_SemanticMetric": used_sem,
                        f"{prefix}_Semantic": sem_sig,
                        f"{prefix}_Cosine": extras.get("cosine"),
                        f"{prefix}_BERTScoreF1": extras.get("bertscore_f1"),
                        f"{prefix}_Accuracy": hybrid,
                        f"{prefix}_Fluency": flu,
                        f"{prefix}_Style": style,
                        f"{prefix}_SQI": sqi,
                        f"{prefix}_LI": li,
                        f"{prefix}_GateA_SemanticOK": "OK" if gateA_pass else "",
                        f"{prefix}_GateB_Flag": "Review" if gateB_flag else "",
                        f"{prefix}_Errors": ", ".join(sorted(set(errs_raw))) if errs_raw else "None",
                        f"{prefix}_ErrorsNorm": ",".join(errs_norm) if errs_norm else "",
                    }
                )

            else:
                errs_raw = []
                if flu < 3:
                    errs_raw.append("Fluency/Grammar")
                if style < 3:
                    errs_raw.append("Style")
                errs_norm = normalize_labels(errs_raw)
                row_result.update(
                    {
                        f"{prefix}_Fluency": flu,
                        f"{prefix}_Style": style,
                        f"{prefix}_Cosine": None,
                        f"{prefix}_Lexical": None,
                        f"{prefix}_Semantic": None,
                        f"{prefix}_Accuracy": None,
                        f"{prefix}_Errors": "None" if not errs_raw else ", ".join(sorted(set(errs_raw))),
                        f"{prefix}_ErrorsNorm": ",".join(errs_norm) if errs_norm else "",
                    }
                )

        results.append(row_result)
        progress.progress((idx + 1) / max(1, len(df)))

    progress.empty()
    res_df = pd.DataFrame(results)

    meta_rows = []
    if meta_mode != "None":
        for i in range(len(res_df)):
            meta_rows.append(
                {
                    "Language": per_row_meta[i]["Language"],
                    "Domain": per_row_meta[i]["Domain"],
                    "Genre": per_row_meta[i]["Genre"],
                }
            )
    meta_df = pd.DataFrame(meta_rows) if meta_rows else pd.DataFrame()

    if mode == "Reference-based":
        acc_cols = [f"{unit['prefix']}_Accuracy" for unit in analysis_units if f"{unit['prefix']}_Accuracy" in res_df.columns]
        if acc_cols:
            best_series = res_df[acc_cols].idxmax(axis=1)
            res_df["Best_Translation"] = best_series.str.replace("_Accuracy", "", regex=False)

    st.dataframe(res_df.head(20))

    lex_vals = collect_metric_series(res_df, "_Lexical")
    flu_vals = collect_metric_series(res_df, "_Fluency")
    sty_vals = collect_metric_series(res_df, "_Style")
    acc_vals = collect_metric_series(res_df, "_Accuracy")

    default_lex_low = 0.15
    default_fluency_low = 3.0
    default_style_low = 3.0
    default_hybrid_low = 0.50

    lex_low = max(default_lex_low, safe_q(lex_vals, 0.10, default_lex_low))
    flu_low = max(default_fluency_low, safe_q(flu_vals, 0.10, default_fluency_low))
    style_low = max(default_style_low, safe_q(sty_vals, 0.10, default_style_low))
    acc_low = safe_q(acc_vals, 0.20, default_hybrid_low)

    st.markdown(
        f"**Calibrated Low Cutoffs:** Lexical < `{lex_low:.2f}`, Fluency < `{flu_low:.2f}`, "
        f"Style < `{style_low:.2f}`, Hybrid Accuracy < `{acc_low:.2f}`"
    )

    st.subheader("Low-Quality Translation Flags")
    flagged_cols = []

    if mode == "Reference-based" and source_col:
        for i in range(len(res_df)):
            acc_thr_adapt, lex_thr_adapt = dynamic_thresholds(ref_lens[i])
            for unit in analysis_units:
                prefix = unit["prefix"]
                acc_col = f"{prefix}_Accuracy"
                lex_col = f"{prefix}_Lexical"
                flu_col = f"{prefix}_Fluency"
                sty_col = f"{prefix}_Style"

                if acc_col in res_df.columns:
                    flag_g = f"{prefix}_Low_Hybrid_Flag_Global"
                    flag_a = f"{prefix}_Low_Hybrid_Flag_Adaptive"
                    val = res_df.loc[i, acc_col]
                    res_df.loc[i, flag_g] = "Review" if pd.notna(val) and val < acc_low else ""
                    res_df.loc[i, flag_a] = "Review" if pd.notna(val) and val < acc_thr_adapt else ""
                    flagged_cols.extend([flag_g, flag_a])

                if lex_col in res_df.columns:
                    flag_g = f"{prefix}_Low_Lexical_Flag_Global"
                    flag_a = f"{prefix}_Low_Lexical_Flag_Adaptive"
                    val = res_df.loc[i, lex_col]
                    res_df.loc[i, flag_g] = "Review" if pd.notna(val) and val < lex_low else ""
                    res_df.loc[i, flag_a] = "Review" if pd.notna(val) and val < lex_thr_adapt else ""
                    flagged_cols.extend([flag_g, flag_a])

                if flu_col in res_df.columns:
                    flag_g = f"{prefix}_Low_Fluency_Flag_Global"
                    val = res_df.loc[i, flu_col]
                    res_df.loc[i, flag_g] = "Review" if pd.notna(val) and val < flu_low else ""
                    flagged_cols.append(flag_g)

                if sty_col in res_df.columns:
                    flag_g = f"{prefix}_Low_Style_Flag_Global"
                    val = res_df.loc[i, sty_col]
                    res_df.loc[i, flag_g] = "Review" if pd.notna(val) and val < style_low else ""
                    flagged_cols.append(flag_g)
    else:
        for col in list(res_df.columns):
            if col.endswith("_Fluency"):
                flag_col = col.replace("_Fluency", "_Low_Fluency_Flag_Global")
                res_df[flag_col] = res_df[col].apply(lambda x: "Review" if pd.notna(x) and x < flu_low else "")
                flagged_cols.append(flag_col)
            if col.endswith("_Style"):
                flag_col = col.replace("_Style", "_Low_Style_Flag_Global")
                res_df[flag_col] = res_df[col].apply(lambda x: "Review" if pd.notna(x) and x < style_low else "")
                flagged_cols.append(flag_col)

    if flagged_cols:
        st.dataframe(res_df[sorted(set(flagged_cols))].head(20))

    st.subheader("Triage Summary")
    if mode == "Reference-based":
        tri_rows = []
        for unit in analysis_units:
            prefix = unit["prefix"]
            ga_col = f"{prefix}_GateA_SemanticOK"
            gb_col = f"{prefix}_GateB_Flag"
            err_col = f"{prefix}_ErrorsNorm"
            ga_ok_pct = float((res_df[ga_col] == "OK").mean() * 100) if ga_col in res_df.columns else np.nan
            gb_flag_pct = float((res_df[gb_col] == "Review").mean() * 100) if gb_col in res_df.columns else np.nan
            paraphrase_cases = int(res_df[err_col].astype(str).str.contains("paraphrase").sum()) if err_col in res_df.columns else 0
            drift_cases = int(res_df[err_col].astype(str).str.contains("meaning_drift|semantic").sum()) if err_col in res_df.columns else 0
            tri_rows.append(
                {
                    "Student": unit["label"],
                    "GateA_OK_%": ga_ok_pct,
                    "GateB_Flag_%": gb_flag_pct,
                    "Paraphrase_cases": paraphrase_cases,
                    "MeaningDrift_cases": drift_cases,
                }
            )
        st.dataframe(pd.DataFrame(tri_rows))

    st.subheader("Disagreements & Teaching Cases")
    disagree_examples = []
    max_examples = 12
    for i in range(len(res_df)):
        for unit in analysis_units:
            prefix = unit["prefix"]
            if f"{prefix}_Semantic" not in res_df.columns:
                continue
            sem_val = res_df.loc[i, f"{prefix}_Semantic"]
            lex_val = res_df.loc[i, f"{prefix}_Lexical"] if f"{prefix}_Lexical" in res_df.columns else np.nan
            cos_val = res_df.loc[i, f"{prefix}_Cosine"] if f"{prefix}_Cosine" in res_df.columns else np.nan
            bs_val = res_df.loc[i, f"{prefix}_BERTScoreF1"] if f"{prefix}_BERTScoreF1" in res_df.columns else np.nan

            cond_paraphrase_like = pd.notna(sem_val) and pd.notna(lex_val) and sem_val >= paraphrase_sem_hi and lex_val < low_lexical_for_paraphrase
            cond_metric_disagree = pd.notna(cos_val) and pd.notna(bs_val) and abs(cos_val - bs_val) > consistency_tolerance
            if cond_paraphrase_like or cond_metric_disagree:
                disagree_examples.append(
                    {
                        "Row": i,
                        "Student": unit["label"],
                        "Semantic": round(float(sem_val), 3) if pd.notna(sem_val) else np.nan,
                        "Lexical": round(float(lex_val), 3) if pd.notna(lex_val) else np.nan,
                        "Cosine": round(float(cos_val), 3) if pd.notna(cos_val) else np.nan,
                        "BERTScoreF1": round(float(bs_val), 3) if pd.notna(bs_val) else np.nan,
                        "Source": str(df.loc[i, unit["reference_col"]])[:220] if unit.get("reference_col") else "",
                        "Translation": str(df.loc[i, unit["translation_col"]])[:220],
                    }
                )
            if len(disagree_examples) >= max_examples:
                break
        if len(disagree_examples) >= max_examples:
            break
    if disagree_examples:
        st.caption("Examples to guide reviewers - useful for spotting valid paraphrases and metric disagreements.")
        st.dataframe(pd.DataFrame(disagree_examples))

    st.subheader("Examples & Issue Mining")
    top_n = st.slider("How many examples per issue?", 3, 50, 10, 1)
    threshold_pack = {"LexLow": lex_low, "AccLow": acc_low, "FluLow": flu_low, "StyleLow": style_low}
    ex_dict = gather_issue_examples(
        res_df=res_df,
        df_orig=df,
        analysis_units=analysis_units,
        mode=mode,
        top_n=top_n,
        thresholds=threshold_pack,
        consistency_tolerance=consistency_tolerance,
        paraphrase_sem_hi=paraphrase_sem_hi,
        low_lexical_for_paraphrase=low_lexical_for_paraphrase,
        drift_sem_lo=drift_sem_lo,
    )

    tabs = st.tabs(list(ex_dict.keys()))
    for tab, (label, exdf) in zip(tabs, ex_dict.items()):
        with tab:
            if exdf.empty:
                st.info(f"No examples for **{label}** with current thresholds.")
            else:
                st.dataframe(exdf)
                c1, c2 = st.columns(2)
                with c1:
                    csv = exdf.to_csv(index=False).encode("utf-8-sig")
                    st.download_button(
                        f"Download {label} (CSV)",
                        csv,
                        f"examples_{label.replace(' ', '_').replace('/', '-').lower()}.csv",
                        "text/csv",
                    )
                with c2:
                    try:
                        docx_bytes = examples_to_docx(exdf, title=f"{label} Examples")
                        st.download_button(
                            f"Download {label} (DOCX)",
                            docx_bytes,
                            f"examples_{label.replace(' ', '_').replace('/', '-').lower()}.docx",
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        )
                    except Exception as e:
                        st.warning(f"DOCX export skipped: {e}")

            if mode in {"Reference-based", "Pairwise Comparison"} and not exdf.empty:
                with st.expander(f"Inline diffs for {label}", expanded=False):
                    for _, r in exdf.iterrows():
                        src = str(r.get("Source", ""))
                        hyp = str(r.get("Translation", ""))
                        st.markdown(token_diff(src, hyp), unsafe_allow_html=True)
                        st.markdown("<hr/>", unsafe_allow_html=True)

    non_empty = [(k, len(v)) for k, v in ex_dict.items() if isinstance(v, pd.DataFrame) and not v.empty and k != "Combined"]
    if non_empty:
        worst_label = sorted(non_empty, key=lambda x: x[1], reverse=True)[0][0]
        st.caption(f"Most populous bucket right now: **{worst_label}**")

    st.subheader("Outlier Monitoring")
    lex_out_pct = iqr_outlier_share(collect_metric_series(res_df, "_Lexical"))
    acc_out_pct = iqr_outlier_share(collect_metric_series(res_df, "_Accuracy"))
    st.write({"Lexical_outliers_%": lex_out_pct, "Hybrid_outliers_%": acc_out_pct})

    if "outlier_history" not in st.session_state:
        st.session_state.outlier_history = []
    st.session_state.outlier_history.append({"Lexical": lex_out_pct, "Hybrid": acc_out_pct, "n": len(res_df)})
    st.session_state.outlier_history = st.session_state.outlier_history[-50:]
    hist_df = pd.DataFrame(st.session_state.outlier_history)
    if not hist_df.empty:
        st.line_chart(hist_df[["Lexical", "Hybrid"]])

    if mode == "Reference-based":
        st.subheader("Per-Sentence Similarity Heatmaps")

        def style_heatmap(df_num):
            try:
                return df_num.style.background_gradient(cmap="YlGnBu").format("{:.2f}")
            except Exception:
                return df_num

        metric_map = {
            "Lexical": [c for c in res_df.columns if c.endswith("_Lexical")],
            "Semantic": [c for c in res_df.columns if c.endswith("_Semantic")],
            "Cosine": [c for c in res_df.columns if c.endswith("_Cosine")],
            "Accuracy": [c for c in res_df.columns if c.endswith("_Accuracy")],
        }
        for metric, cols in metric_map.items():
            if cols:
                st.markdown(f"**{metric} per Sentence**")
                st.dataframe(style_heatmap(res_df[cols]))

    st.subheader("Dashboard Summary")
    metrics = ["Fluency", "Style"] if mode == "Standalone Student Assessment" else ["Lexical", "Semantic", "Cosine", "Accuracy", "Fluency", "Style", "SQI", "LI"]
    for metric in metrics:
        metric_cols = [c for c in res_df.columns if c.endswith(metric)]
        if metric_cols:
            show_metric_boxplot(res_df, metric_cols, metric)

    st.subheader("Per-Domain Diagnostics")
    if not meta_df.empty and meta_df.shape[0] == res_df.shape[0]:
        long_rows = []
        for i in range(len(res_df)):
            meta_row = meta_df.iloc[i].to_dict()
            for unit in analysis_units:
                prefix = unit["prefix"]
                long_rows.append(
                    {
                        **meta_row,
                        "Student": unit["label"],
                        "Accuracy": res_df.get(f"{prefix}_Accuracy", pd.Series([np.nan] * len(res_df))).iloc[i],
                        "Fluency": res_df.get(f"{prefix}_Fluency", pd.Series([np.nan] * len(res_df))).iloc[i],
                        "Style": res_df.get(f"{prefix}_Style", pd.Series([np.nan] * len(res_df))).iloc[i],
                        "Lexical": res_df.get(f"{prefix}_Lexical", pd.Series([np.nan] * len(res_df))).iloc[i],
                        "Semantic": res_df.get(f"{prefix}_Semantic", pd.Series([np.nan] * len(res_df))).iloc[i],
                        "ErrorsNorm": res_df.get(f"{prefix}_ErrorsNorm", pd.Series([""] * len(res_df))).iloc[i],
                    }
                )
        long_df = pd.DataFrame(long_rows)

        if not long_df.empty:
            group_keys = [k for k in ["Domain", "Language", "Genre"] if k in long_df.columns]
            if not group_keys:
                st.caption("No metadata provided; skipping per-domain aggregates.")
            else:
                def top_labels(series: pd.Series):
                    labels = ",".join(series.dropna().astype(str)).split(",")
                    labels = [x for x in labels if x]
                    if not labels:
                        return ""
                    s = pd.Series(labels).value_counts(normalize=True)
                    pairs = [f"{lab}:{share:.0%}" for lab, share in s.head(3).items()]
                    return ", ".join(pairs)

                agg = long_df.groupby(group_keys).agg(
                    Count=("Accuracy", "count"),
                    Accuracy_Mean=("Accuracy", "mean"),
                    Accuracy_Std=("Accuracy", "std"),
                    Fluency_Mean=("Fluency", "mean"),
                    Fluency_Std=("Fluency", "std"),
                    Style_Mean=("Style", "mean"),
                    Lexical_Mean=("Lexical", "mean"),
                    Semantic_Mean=("Semantic", "mean"),
                    Top_Labels=("ErrorsNorm", top_labels),
                ).reset_index()
                st.dataframe(agg)
    else:
        st.caption("No metadata or shape mismatch; per-domain diagnostics skipped.")

    st.subheader("Export Cleaned Results")
    preferred_order = []
    ordered_metrics = [
        "Accuracy",
        "Lexical",
        "Semantic",
        "Cosine",
        "BERTScoreF1",
        "Fluency",
        "Style",
        "SQI",
        "LI",
        "GateA_SemanticOK",
        "GateB_Flag",
        "Errors",
        "ErrorsNorm",
        "SemanticMetric",
        "LexicalMetric",
        "Low_Hybrid_Flag_Global",
        "Low_Hybrid_Flag_Adaptive",
        "Low_Lexical_Flag_Global",
        "Low_Lexical_Flag_Adaptive",
        "Low_Fluency_Flag_Global",
        "Low_Style_Flag_Global",
    ]

    for unit in analysis_units:
        prefix = unit["prefix"]
        for metric in ordered_metrics:
            col = f"{prefix}_{metric}"
            if col in res_df.columns:
                preferred_order.append(col)

    if "Best_Translation" in res_df.columns:
        preferred_order = ["Best_Translation"] + preferred_order

    other_cols = [c for c in res_df.columns if c not in preferred_order]
    ordered_cols = preferred_order + other_cols

    export_df = res_df[ordered_cols].copy()
    if not meta_df.empty:
        export_df = pd.concat([meta_df, export_df], axis=1)

    export_df.columns = humanize_export_columns(export_df.columns)
    export_df = export_df.applymap(lambda x: round(x, 3) if isinstance(x, (float, int)) else x)

    st.dataframe(export_df.head(20))
    csv = export_df.to_csv(index=False, encoding="utf-8-sig").encode("utf-8-sig")
    st.download_button(
        "Download Full Analysis Results (Clean CSV)",
        csv,
        "translation_analysis_clean.csv",
        "text/csv",
    )

    with st.expander("Thresholds (Quantile-based) Details", expanded=False):
        st.write(
            {
                "Lexical_Low": lex_low,
                "Fluency_Low": flu_low,
                "Style_Low": style_low,
                "Hybrid_Low": acc_low,
            }
        )

except Exception as e:
    st.error(f"Error: {e}")
